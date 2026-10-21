#!/usr/bin/env python3
import argparse
import json
import re
import unicodedata
import urllib.request
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from difflib import SequenceMatcher
from pathlib import Path

from openpyxl import load_workbook


LEGAL_SUFFIXES = (
    "股份有限责任公司",
    "集团股份有限公司",
    "股份有限公司",
    "集团有限公司",
    "有限责任公司",
    "责任有限公司",
    "控股有限公司",
    "有限公司",
    "集团",
    "公司",
)
CITY_PREFIXES = {
    "深圳市": "深圳",
    "东莞市": "东莞",
    "惠州市": "惠州",
    "广州市": "广州",
    "佛山市": "佛山",
    "中山市": "中山",
    "珠海市": "珠海",
    "江门市": "江门",
}
PROVINCE_ZH = {
    "Guangdong": "广东省",
    "Hunan": "湖南省",
    "Anhui": "安徽省",
    "Jiangxi": "江西省",
    "Zhejiang": "浙江省",
}
CITY_ZH = {
    "Shenzhen": "深圳市",
    "Dongguan": "东莞市",
    "Huizhou": "惠州市",
    "Guangzhou": "广州市",
    "Hengyang": "衡阳市",
    "Hefei": "合肥市",
    "Nanchang": "南昌市",
    "Zhaoqing": "肇庆市",
    "Zhoushan": "舟山市",
    "Zhuhai": "珠海市",
}
VALID_PHASES = {"第1期", "第2期", "第3期"}


def normalize_name(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"[\s\u200b\ufeff·•,，。．、()（）]+", "", text).strip()


def core_name(value: str) -> str:
    text = normalize_name(value)
    for old, new in CITY_PREFIXES.items():
        if text.startswith(old):
            text = new + text[len(old) :]
            break
    changed = True
    while changed:
        changed = False
        for suffix in LEGAL_SUFFIXES:
            if text.endswith(suffix) and len(text) > len(suffix) + 2:
                text = text[: -len(suffix)]
                changed = True
                break
    return text


def trigrams(text: str) -> set[str]:
    if len(text) < 3:
        return {text}
    return {text[index : index + 3] for index in range(len(text) - 2)}


def edit_distance(left: str, right: str, max_distance: int = 2) -> int:
    if abs(len(left) - len(right)) > max_distance:
        return max_distance + 1
    previous = list(range(len(right) + 1))
    for index, left_char in enumerate(left, 1):
        current = [index]
        row_minimum = index
        for right_index, right_char in enumerate(right, 1):
            value = min(
                current[-1] + 1,
                previous[right_index] + 1,
                previous[right_index - 1] + (left_char != right_char),
            )
            current.append(value)
            row_minimum = min(row_minimum, value)
        if row_minimum > max_distance:
            return max_distance + 1
        previous = current
    return previous[-1]


def split_values(value: str) -> list[str]:
    values = []
    for item in re.split(r"[、;；\n]+", str(value or "")):
        item = item.strip(" ，,")
        if item and item not in values:
            values.append(item)
    return values


def phase_values(value: str) -> list[str]:
    phases = []
    for phase in re.findall(r"第[123]期", str(value or "")):
        if phase not in phases:
            phases.append(phase)
    return phases


def load_canton(path: str) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["參展商名單"]
    rows = []
    for values in worksheet.iter_rows(min_row=4, values_only=True):
        if not values[1]:
            continue
        rows.append(
            {
                "supplier_no": str(values[0] or "").strip(),
                "name": str(values[1]).strip(),
                "phase": phase_values(values[2]),
                "section": split_values(values[3]),
                "booth": split_values(values[4]),
            }
        )
    return rows


def load_hedge(path: str) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["原始数据"]
    rows_by_name = {}
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not values[0]:
            continue
        name = str(values[0]).strip()
        rows_by_name.setdefault(
            normalize_name(name),
            {"name": name, "org_code": str(values[1] or "").strip()},
        )
    return list(rows_by_name.values())


def match_ab(canton: list[dict], hedge: list[dict]) -> list[dict]:
    hedge_normalized = {normalize_name(row["name"]): row for row in hedge}
    hedge_core = defaultdict(list)
    gram_index = defaultdict(set)
    for index, row in enumerate(hedge):
        core = core_name(row["name"])
        hedge_core[core].append(index)
        if len(core) >= 5:
            for gram in trigrams(core):
                gram_index[gram].add(index)

    matches = []
    for exhibitor in canton:
        normalized = normalize_name(exhibitor["name"])
        if normalized in hedge_normalized:
            matches.append(
                {
                    **exhibitor,
                    "match_grade": "A",
                    "match_reason": "完整名称一致",
                    "hedge_name": hedge_normalized[normalized]["name"],
                }
            )
            continue

        core = core_name(exhibitor["name"])
        exact_core_candidates = hedge_core.get(core, [])
        if len(core) >= 4 and len(exact_core_candidates) == 1:
            hedge_row = hedge[exact_core_candidates[0]]
            matches.append(
                {
                    **exhibitor,
                    "match_grade": "B",
                    "match_reason": "核心名称一致，法律形式或地名写法不同",
                    "hedge_name": hedge_row["name"],
                }
            )
            continue

        if len(core) < 6:
            continue
        candidates = Counter()
        for gram in trigrams(core):
            for index in gram_index.get(gram, ()):
                candidates[index] += 1
        scored = []
        for index, shared in candidates.most_common(80):
            candidate_core = core_name(hedge[index]["name"])
            if abs(len(core) - len(candidate_core)) > 2:
                continue
            distance = edit_distance(core, candidate_core, 2)
            if distance > 2:
                continue
            ratio = SequenceMatcher(None, core, candidate_core).ratio()
            if ratio < 0.92:
                continue
            scored.append((ratio, -distance, shared, index))
        scored.sort(reverse=True)
        if not scored:
            continue
        second_ratio = scored[1][0] if len(scored) > 1 else 0
        if scored[0][0] - second_ratio < 0.03:
            continue
        best = scored[0]
        matches.append(
            {
                **exhibitor,
                "match_grade": "B",
                "match_reason": f"核心名称近似，编辑距离{-best[1]}",
                "hedge_name": hedge[best[3]]["name"],
            }
        )
    return matches


def fetch_supplier_detail(supplier_no: str, cache_dir: str) -> dict:
    cache_path = Path(cache_dir) / f"{supplier_no}.json"
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    if cache_path.exists():
        return json.loads(cache_path.read_text(encoding="utf-8"))

    url = (
        "https://www.cantonfair.org.cn/appbuyerapi/supplier/detail?supplierNo="
        + supplier_no
    )
    request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    last_error = None
    for _ in range(3):
        try:
            with urllib.request.urlopen(request, timeout=25) as response:
                payload = json.load(response)
            result = payload.get("result") or {}
            if not result:
                raise ValueError(f"empty supplier result for {supplier_no}")
            cache_path.write_text(
                json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            return result
        except Exception as error:
            last_error = error
    raise RuntimeError(f"failed supplier detail {supplier_no}: {last_error}")


def localize_province(value):
    if not value:
        return None
    return PROVINCE_ZH.get(value, value if str(value).endswith("省") else f"{value}省")


def localize_city(value):
    if not value:
        return None
    return CITY_ZH.get(value, value if str(value).endswith("市") else f"{value}市")


def enrich_rows(rows: list[dict], cache_dir: str, workers: int = 12) -> list[dict]:
    details = {}
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(fetch_supplier_detail, row["supplier_no"], cache_dir): row[
                "supplier_no"
            ]
            for row in rows
        }
        for future in as_completed(futures):
            supplier_no = futures[future]
            details[supplier_no] = future.result()

    enriched = []
    for row in rows:
        result = details[row["supplier_no"]]
        introduction = result.get("companyIntroduction") or {}
        contact = result.get("companyContact") or {}
        enriched.append(
            {
                **row,
                "province": localize_province(introduction.get("provinceName")),
                "city": localize_city(introduction.get("cityName")),
                "address": introduction.get("companyAddress")
                or contact.get("companyAddress")
                or None,
                "capital": introduction.get("registeredCapital") or None,
                "official_url": (
                    "https://365.cantonfair.org.cn/zh-CN/shops?id="
                    + row["supplier_no"]
                ),
                "data_status": "广交会展商资料，未经工商核验",
            }
        )
    return enriched


def allocate_ids(base_rows: list[dict], canton_rows: list[dict]) -> list[dict]:
    base_by_name = {normalize_name(row.get("name")): row for row in base_rows}
    next_id = max((int(row["id"]) for row in base_rows), default=-1) + 1
    allocated = []
    for row in canton_rows:
        existing = base_by_name.get(normalize_name(row["name"]))
        allocated.append({**row, "id": int(existing["id"]) if existing else next_id})
        if not existing:
            next_id += 1
    return allocated


def validate_rows(
    rows: list[dict], expected_total: int = 300, expected_a: int = 294, expected_b: int = 6
) -> None:
    if len(rows) != expected_total:
        raise ValueError(f"expected {expected_total} rows, got {len(rows)}")
    grade_counts = Counter(row.get("match_grade") for row in rows)
    if grade_counts != Counter({"A": expected_a, "B": expected_b}):
        raise ValueError(f"unexpected grade counts: {dict(grade_counts)}")
    ids = [row.get("id") for row in rows]
    if len(ids) != len(set(ids)):
        raise ValueError("duplicate Hunter Hub ids")
    names = [normalize_name(row.get("name")) for row in rows]
    if len(names) != len(set(names)):
        raise ValueError("duplicate normalized company names")
    for row in rows:
        if not row.get("name") or not row.get("supplier_no"):
            raise ValueError("missing company name or supplier number")
        phases = row.get("phase") or []
        if not phases or any(phase not in VALID_PHASES for phase in phases):
            raise ValueError(f"invalid phase for {row.get('name')}: {phases}")
        if not str(row.get("official_url", "")).startswith(
            "https://365.cantonfair.org.cn/"
        ):
            raise ValueError(f"invalid official URL for {row.get('name')}")


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument("--canton-workbook", required=True)
    parser.add_argument("--hedge-workbook", required=True)
    parser.add_argument("--base-data", required=True)
    parser.add_argument("--cache-dir", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--workers", type=int, default=16)
    return parser.parse_args()


def main():
    args = parse_args()
    canton_rows = load_canton(args.canton_workbook)
    hedge_rows = load_hedge(args.hedge_workbook)
    matched = match_ab(canton_rows, hedge_rows)
    enriched = enrich_rows(matched, args.cache_dir, args.workers)
    with open(args.base_data, encoding="utf-8") as handle:
        base_rows = json.load(handle)
    allocated = allocate_ids(base_rows, enriched)
    validate_rows(allocated)

    output_path = Path(args.output)
    temporary_path = output_path.with_suffix(output_path.suffix + ".tmp")
    temporary_path.write_text(
        json.dumps(allocated, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    temporary_path.replace(output_path)

    existing_ids = {row["id"] for row in base_rows}
    existing_count = sum(row["id"] in existing_ids for row in allocated)
    grades = Counter(row["match_grade"] for row in allocated)
    print(
        f"A={grades['A']} B={grades['B']} total={len(allocated)} "
        f"existing={existing_count} new={len(allocated) - existing_count}"
    )


if __name__ == "__main__":
    main()
